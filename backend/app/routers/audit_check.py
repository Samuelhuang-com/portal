"""
稽核檢查（Audit Check）— FastAPI Router
Prefix: /api/v1/audit-check

中文名稱：稽核檢查 ／ 英文名稱：Audit Check
前端路由：/audit-check/*（手機版 /m/audit-check）
對應 Excel：財_3系統建置稽核-202609執行.xlsx
規格：docs/SPEC_audit_check.md

權限（CLAUDE.md §11）
    audit_check_view   檢視稽核單與統計
    audit_check_edit   建立期別／稽核單、填寫評語與覆核
    audit_check_admin  檢查項主檔、判定類型設定、刪除

⚠️ 所有端點一律掛 require_permission，不得只掛 get_current_user（§11.6）。
⚠️ 本模組非 Ragic 同步模組，sync_tool.py / RagicConnections.tsx 一律不動。
"""
from __future__ import annotations

import io
import re
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.time import twnow
from app.dependencies import require_permission
from app.models.audit_check import (
    AuditCell, AuditItem, AuditPeriod, AuditResultType,
    AuditReview, AuditSheet, AuditSheetDepartment, AuditSheetItem,
)
from app.models.reference_data import Company, RefDepartment
from app.models.user import User
from app.schemas.audit_check import (
    CellBulkUpsert, CellUpsert, DeficiencyUpsert, ItemCreate, ItemOut, ItemUpdate,
    PeriodCreate, PeriodOut, PeriodUpdate, ResultTypeCreate, ResultTypeOut,
    ResultTypeUpdate, ReviewUpsert, SheetCreate, SheetDetail, SheetLayoutUpdate,
    SheetUpdate, StatisticsOut,
)
from app.services import audit_check_service as svc

router = APIRouter()

VIEW = "audit_check_view"
EDIT = "audit_check_edit"
ADMIN = "audit_check_admin"


def _conflict(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_409_CONFLICT, detail=detail)


def _not_found(detail: str = "查無資料") -> HTTPException:
    return HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=detail)


# ════════════════════════════════════════════════════════════════════════════
# 判定類型（使用者自訂字色語意）
# ════════════════════════════════════════════════════════════════════════════
@router.get("/result-types", response_model=List[ResultTypeOut])
def list_result_types(
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    return svc.get_result_types(db)


@router.post("/result-types", response_model=ResultTypeOut, status_code=201)
def create_result_type(
    payload: ResultTypeCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    svc.seed_result_types(db)
    if db.query(AuditResultType).filter(AuditResultType.code == payload.code).first():
        raise _conflict("代碼已存在")
    rt = AuditResultType(**payload.model_dump(), is_system=False, is_active=True)
    db.add(rt)
    if payload.is_default:
        db.flush()
        db.query(AuditResultType).filter(AuditResultType.id != rt.id).update({"is_default": False})
    db.commit()
    db.refresh(rt)
    return rt


@router.put("/result-types/{rt_id}", response_model=ResultTypeOut)
def update_result_type(
    rt_id: int,
    payload: ResultTypeUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    rt = db.query(AuditResultType).filter(AuditResultType.id == rt_id).first()
    if rt is None:
        raise _not_found("查無此判定類型")
    data = payload.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(rt, k, v)
    if data.get("is_default"):
        db.flush()
        db.query(AuditResultType).filter(AuditResultType.id != rt.id).update({"is_default": False})
    db.commit()
    db.refresh(rt)
    return rt


@router.patch("/result-types/{rt_id}/toggle", response_model=ResultTypeOut)
def toggle_result_type(
    rt_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    rt = db.query(AuditResultType).filter(AuditResultType.id == rt_id).first()
    if rt is None:
        raise _not_found("查無此判定類型")
    if rt.is_active and rt.is_default:
        raise _conflict("預設判定類型不可停用，請先把預設改給其他類型")
    rt.is_active = not rt.is_active
    db.commit()
    db.refresh(rt)
    return rt


@router.delete("/result-types/{rt_id}", status_code=204)
def delete_result_type(
    rt_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    rt = db.query(AuditResultType).filter(AuditResultType.id == rt_id).first()
    if rt is None:
        raise _not_found("查無此判定類型")
    if rt.is_system:
        raise _conflict("系統內建的判定類型不可刪除，可改名、改色或停用")
    used = db.query(AuditCell.id).filter(AuditCell.result_code == rt.code).first()
    if used is not None:
        raise _conflict("此判定類型已被稽核資料使用，不可刪除，請改用停用")
    db.delete(rt)
    db.commit()
    return Response(status_code=204)


# ════════════════════════════════════════════════════════════════════════════
# 檢查項主檔
# ════════════════════════════════════════════════════════════════════════════
def _item_tree(db: Session, include_inactive: bool) -> List[dict]:
    q = db.query(AuditItem)
    if not include_inactive:
        q = q.filter(AuditItem.is_active.is_(True))
    rows = q.order_by(AuditItem.sort_order, AuditItem.id).all()
    used = svc.items_in_use(db)

    def to_dict(it: AuditItem) -> dict:
        return {
            "id": it.id,
            "parent_id": it.parent_id,
            "name": it.name,
            "description": it.description,
            "sort_order": it.sort_order,
            "is_active": it.is_active,
            "in_use": it.id in used,
            "children": [],
        }

    nodes = {it.id: to_dict(it) for it in rows}
    tree: List[dict] = []
    for it in rows:
        node = nodes[it.id]
        if it.parent_id and it.parent_id in nodes:
            nodes[it.parent_id]["children"].append(node)
        elif it.parent_id is None:
            tree.append(node)
    return tree


@router.get("/items", response_model=List[ItemOut])
def list_items(
    include_inactive: bool = Query(False),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    return _item_tree(db, include_inactive)


@router.post("/items", response_model=ItemOut, status_code=201)
def create_item(
    payload: ItemCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    if payload.parent_id is not None:
        parent = db.query(AuditItem).filter(AuditItem.id == payload.parent_id).first()
        if parent is None:
            raise _not_found("查無上層大項")
        if parent.parent_id is not None:
            raise _conflict("檢查項只有兩階，不能掛在子項底下")
    item = AuditItem(**payload.model_dump())
    db.add(item)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise _conflict("同一層已有相同名稱的檢查項")
    db.refresh(item)
    return {
        "id": item.id, "parent_id": item.parent_id, "name": item.name,
        "description": item.description, "sort_order": item.sort_order,
        "is_active": item.is_active, "in_use": False, "children": [],
    }


@router.put("/items/{item_id}", response_model=ItemOut)
def update_item(
    item_id: int,
    payload: ItemUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    item = db.query(AuditItem).filter(AuditItem.id == item_id).first()
    if item is None:
        raise _not_found("查無此檢查項")
    data = payload.model_dump(exclude_unset=True)
    # 使用者裁示：已被任一期稽核單引用即鎖定，不可改名（保護歷史資料）
    if "name" in data and data["name"] != item.name and svc.item_in_use(db, item_id):
        raise _conflict("此檢查項已被稽核單引用，不可改名；如不再使用請改為「停用」")
    for k, v in data.items():
        setattr(item, k, v)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise _conflict("同一層已有相同名稱的檢查項")
    db.refresh(item)
    return {
        "id": item.id, "parent_id": item.parent_id, "name": item.name,
        "description": item.description, "sort_order": item.sort_order,
        "is_active": item.is_active, "in_use": svc.item_in_use(db, item_id), "children": [],
    }


@router.patch("/items/{item_id}/toggle", response_model=ItemOut)
def toggle_item(
    item_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    item = db.query(AuditItem).filter(AuditItem.id == item_id).first()
    if item is None:
        raise _not_found("查無此檢查項")
    item.is_active = not item.is_active
    # 停用大項時連同子項一起停用，避免出現「孤兒子項」
    if item.parent_id is None:
        db.query(AuditItem).filter(AuditItem.parent_id == item.id).update({"is_active": item.is_active})
    db.commit()
    db.refresh(item)
    return {
        "id": item.id, "parent_id": item.parent_id, "name": item.name,
        "description": item.description, "sort_order": item.sort_order,
        "is_active": item.is_active, "in_use": svc.item_in_use(db, item_id), "children": [],
    }


@router.delete("/items/{item_id}", status_code=204)
def delete_item(
    item_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    item = db.query(AuditItem).filter(AuditItem.id == item_id).first()
    if item is None:
        raise _not_found("查無此檢查項")
    if svc.item_in_use(db, item_id):
        raise _conflict("此檢查項已被稽核單引用，不可刪除；如不再使用請改為「停用」")
    child_ids = [c.id for c in db.query(AuditItem).filter(AuditItem.parent_id == item_id).all()]
    for cid in child_ids:
        if svc.item_in_use(db, cid):
            raise _conflict("底下的子項已被稽核單引用，不可刪除；如不再使用請改為「停用」")
    db.delete(item)
    db.commit()
    return Response(status_code=204)


# ════════════════════════════════════════════════════════════════════════════
# 期別
# ════════════════════════════════════════════════════════════════════════════
def _period_out(db: Session, p: AuditPeriod) -> dict:
    companies = {c.id: c.name for c in db.query(Company).all()}
    sheets = []
    for s in sorted(p.sheets, key=lambda x: x.company_id):
        rate, label = svc.compute_completion(db, s)
        sheets.append({
            "id": s.id,
            "company_id": s.company_id,
            "company_name": companies.get(s.company_id, ""),
            "audited_on": s.audited_on,
            "status": s.status,
            "completion_rate": rate,
            "completion_label": label,
        })
    return {
        "id": p.id, "period": p.period, "title": p.title,
        "goal_major": p.goal_major, "goal_minor": p.goal_minor,
        "note": p.note, "review_counts_in_score": p.review_counts_in_score,
        "status": p.status, "created_at": p.created_at,
        "sheets": sheets,
    }


@router.get("/periods", response_model=List[PeriodOut])
def list_periods(
    year: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    q = db.query(AuditPeriod)
    if year is not None:
        q = q.filter(AuditPeriod.period.like(f"{year}-%"))
    return [_period_out(db, p) for p in q.order_by(AuditPeriod.period.desc()).all()]


@router.post("/periods", response_model=PeriodOut, status_code=201)
def create_period(
    payload: PeriodCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    if db.query(AuditPeriod).filter(AuditPeriod.period == payload.period).first():
        raise _conflict("該期別已存在")
    p = AuditPeriod(**payload.model_dump(), created_by=str(user.id))
    db.add(p)
    db.commit()
    db.refresh(p)
    return _period_out(db, p)


@router.put("/periods/{period_id}", response_model=PeriodOut)
def update_period(
    period_id: int,
    payload: PeriodUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    p = db.query(AuditPeriod).filter(AuditPeriod.id == period_id).first()
    if p is None:
        raise _not_found("查無此期別")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return _period_out(db, p)


@router.delete("/periods/{period_id}", status_code=204)
def delete_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    p = db.query(AuditPeriod).filter(AuditPeriod.id == period_id).first()
    if p is None:
        raise _not_found("查無此期別")
    if p.sheets:
        raise _conflict("此期別底下尚有稽核單，請先刪除稽核單")
    db.delete(p)
    db.commit()
    return Response(status_code=204)


# ════════════════════════════════════════════════════════════════════════════
# 稽核單
# ════════════════════════════════════════════════════════════════════════════
def _get_sheet(db: Session, sheet_id: int) -> AuditSheet:
    s = db.query(AuditSheet).filter(AuditSheet.id == sheet_id).first()
    if s is None:
        raise _not_found("查無此稽核單")
    return s


@router.get("/sheets/{sheet_id}", response_model=SheetDetail)
def get_sheet(
    sheet_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    return svc.build_sheet_detail(db, _get_sheet(db, sheet_id))


@router.post("/sheets", response_model=SheetDetail, status_code=201)
def create_sheet(
    payload: SheetCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    period = db.query(AuditPeriod).filter(AuditPeriod.id == payload.period_id).first()
    if period is None:
        raise _not_found("查無此期別")
    if db.query(Company).filter(Company.id == payload.company_id).first() is None:
        raise _not_found("查無此公司別")
    if db.query(AuditSheet).filter(
        AuditSheet.period_id == payload.period_id,
        AuditSheet.company_id == payload.company_id,
    ).first():
        raise _conflict("該期別的這家公司已經有稽核單了")

    sheet = AuditSheet(
        period_id=payload.period_id,
        company_id=payload.company_id,
        audited_on=payload.audited_on,
        audited_session=payload.audited_session,
        executor_label=payload.executor_label,
        status="draft",
        created_by=str(user.id),
    )
    db.add(sheet)
    db.flush()

    svc.apply_layout(db, sheet, payload.department_ids, payload.items)
    db.flush()
    db.refresh(sheet)
    if payload.carry_over:
        svc.carry_over_reviews(db, sheet)
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


@router.put("/sheets/{sheet_id}", response_model=SheetDetail)
def update_sheet(
    sheet_id: int,
    payload: SheetUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(sheet, k, v)
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


@router.put("/sheets/{sheet_id}/layout", response_model=SheetDetail)
def update_sheet_layout(
    sheet_id: int,
    payload: SheetLayoutUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    svc.apply_layout(db, sheet, payload.department_ids, payload.items)
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


@router.delete("/sheets/{sheet_id}", status_code=204)
def delete_sheet(
    sheet_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    db.delete(_get_sheet(db, sheet_id))
    db.commit()
    return Response(status_code=204)


# ── 格子 upsert ────────────────────────────────────────────────────────────
def _upsert_cell(db: Session, sheet: AuditSheet, item: CellUpsert, default_code: str, valid_codes: set, user_id: str) -> None:
    if item.sheet_item_id not in {si.id for si in sheet.items}:
        raise _not_found("此稽核單沒有這一列檢查項")
    if item.sheet_department_id not in {sd.id for sd in sheet.departments}:
        raise _not_found("此稽核單沒有這一欄部門")

    cell = (
        db.query(AuditCell)
        .filter(
            AuditCell.sheet_item_id == item.sheet_item_id,
            AuditCell.sheet_department_id == item.sheet_department_id,
        )
        .first()
    )
    text = (item.comment or "").strip()
    if not text:
        # 空白 ＝ 該部門本期不查此項，直接刪除不留列（§4.1 分母口徑）
        if cell is not None:
            db.delete(cell)
        return

    code = item.result_code or (cell.result_code if cell else default_code)
    if code not in valid_codes:
        raise _conflict(f"未知的判定類型：{code}")

    if cell is None:
        cell = AuditCell(
            sheet_id=sheet.id,
            sheet_item_id=item.sheet_item_id,
            sheet_department_id=item.sheet_department_id,
        )
        db.add(cell)
    cell.comment = text
    cell.result_code = code
    cell.updated_by = user_id
    cell.updated_at = twnow()


@router.put("/sheets/{sheet_id}/cells", response_model=SheetDetail)
def upsert_cell(
    sheet_id: int,
    payload: CellUpsert,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    by_code, default_code = svc.result_type_maps(db)
    _upsert_cell(db, sheet, payload, default_code, set(by_code.keys()), str(user.id))
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


@router.put("/sheets/{sheet_id}/cells/bulk", response_model=SheetDetail)
def upsert_cells_bulk(
    sheet_id: int,
    payload: CellBulkUpsert,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    by_code, default_code = svc.result_type_maps(db)
    codes = set(by_code.keys())
    for item in payload.cells:
        _upsert_cell(db, sheet, item, default_code, codes, str(user.id))
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


# ── 覆核區 / 缺失覆寫 ──────────────────────────────────────────────────────
@router.put("/sheets/{sheet_id}/reviews/{sheet_department_id}", response_model=SheetDetail)
def upsert_review(
    sheet_id: int,
    sheet_department_id: int,
    payload: ReviewUpsert,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    if sheet_department_id not in {sd.id for sd in sheet.departments}:
        raise _not_found("此稽核單沒有這一欄部門")
    review = (
        db.query(AuditReview)
        .filter(AuditReview.sheet_id == sheet_id, AuditReview.sheet_department_id == sheet_department_id)
        .first()
    )
    if review is None:
        review = AuditReview(sheet_id=sheet_id, sheet_department_id=sheet_department_id)
        db.add(review)
    for k, v in payload.model_dump(exclude_unset=True).items():
        if v is not None or k in ("pending_text", "result_text", "source_period"):
            setattr(review, k, v)
    review.updated_at = twnow()
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


@router.put("/sheets/{sheet_id}/departments/{sheet_department_id}/deficiency", response_model=SheetDetail)
def upsert_deficiency(
    sheet_id: int,
    sheet_department_id: int,
    payload: DeficiencyUpsert,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    sheet = _get_sheet(db, sheet_id)
    sd = next((d for d in sheet.departments if d.id == sheet_department_id), None)
    if sd is None:
        raise _not_found("此稽核單沒有這一欄部門")
    text = (payload.deficiency_override or "").strip()
    sd.deficiency_override = text or None   # 空字串 → 還原成自動彙整
    db.commit()
    db.refresh(sheet)
    return svc.build_sheet_detail(db, sheet)


# ════════════════════════════════════════════════════════════════════════════
# 統計
# ════════════════════════════════════════════════════════════════════════════
@router.get("/statistics", response_model=StatisticsOut)
def statistics(
    year: int = Query(..., ge=2000, le=2999),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    return svc.build_statistics(db, year, company_id)


@router.get("/statistics/flagged")
def flagged_rows(
    period_id: Optional[int] = Query(None),
    company_id: Optional[int] = Query(None),
    department_id: Optional[int] = Query(None),
    only_failing: bool = Query(False, description="True 只回未達標（counts_as_pass=False）"),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """
    缺失／建議清單 — 把所有「列入彙整」的判定類型攤平成清單，
    供跨月追蹤某部門的缺失或建議事項。
    """
    by_code, _default = svc.result_type_maps(db)
    q = db.query(AuditCell).join(AuditSheet, AuditSheet.id == AuditCell.sheet_id)
    if period_id is not None:
        q = q.filter(AuditSheet.period_id == period_id)
    if company_id is not None:
        q = q.filter(AuditSheet.company_id == company_id)

    companies = {c.id: c.name for c in db.query(Company).all()}
    dept_names = {d.id: d.name for d in db.query(RefDepartment).all()}
    sheet_depts = {sd.id: sd for sd in db.query(AuditSheetDepartment).all()}
    sheet_items = {si.id: si for si in db.query(AuditSheetItem).all()}
    items = {i.id: i for i in db.query(AuditItem).all()}
    periods = {p.id: p.period for p in db.query(AuditPeriod).all()}
    sheets = {s.id: s for s in db.query(AuditSheet).all()}

    rows = []
    for c in q.all():
        rt = by_code.get(c.result_code)
        if rt is None or not rt.include_in_summary:
            continue
        if only_failing and rt.counts_as_pass:
            continue
        sd = sheet_depts.get(c.sheet_department_id)
        if sd is None:
            continue
        if department_id is not None and sd.department_id != department_id:
            continue
        si = sheet_items.get(c.sheet_item_id)
        sheet = sheets.get(c.sheet_id)
        rows.append({
            "period": periods.get(sheet.period_id, "") if sheet else "",
            "company_name": companies.get(sheet.company_id, "") if sheet else "",
            "department_id": sd.department_id,
            "department_name": dept_names.get(sd.department_id, ""),
            "item_name": (items.get(si.item_id).name if si and si.item_id in items else ""),
            "display_no": si.display_no if si else "",
            "result_code": c.result_code,
            "result_label": rt.label,
            "result_color": rt.color,
            "comment": c.comment or "",
        })
    rows.sort(key=lambda r: (r["period"], r["company_name"], r["department_name"], r["display_no"]))
    return rows


# ════════════════════════════════════════════════════════════════════════════
# 匯出 Excel
# ════════════════════════════════════════════════════════════════════════════
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(val) -> str:
    """openpyxl 遇到控制字元會丟 IllegalCharacterError → 整支匯出 500。"""
    if val is None:
        return ""
    return _ILLEGAL.sub("", str(val))


def _ascii_slug(text: str) -> str:
    out = re.sub(r"[^A-Za-z0-9._-]+", "_", text or "")
    return out.strip("_") or "export"


def _xlsx_response(wb, filename_cn: str) -> Response:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    encoded = quote(filename_cn, safe="")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{_ascii_slug(filename_cn)}"; '
                f"filename*=UTF-8''{encoded}"
            )
        },
    )


@router.get("/sheets/{sheet_id}/export")
def export_sheet(
    sheet_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    sheet = _get_sheet(db, sheet_id)
    detail = svc.build_sheet_detail(db, sheet)
    colors = {rt.code: (rt.color or "#000000").lstrip("#").upper() for rt in detail["result_types"]}

    wb = Workbook()
    ws = wb.active
    ws.title = _clean(f"{detail['title']}-{detail['period'].replace('-', '')}")[:31]

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    depts = detail["departments"]

    def put(row: int, col: int, value, color: Optional[str] = None):
        cell = ws.cell(row=row, column=col, value=_clean(value))
        cell.font = Font(bold=True, color=("FF" + color) if color else None)
        cell.alignment = wrap
        return cell

    r = 1
    put(r, 1, detail["company_name"])
    put(r, 2, detail["executor_label"] or "")
    put(r, 3, detail["audited_label"], "FF0000")
    for i, d in enumerate(depts):
        put(r, 4 + i, d["name"])
    r += 1

    scores = {s["sheet_department_id"]: s for s in detail["scores"]}
    for label, key in (("稽核子項數", "sub_count"), ("達標項數", "pass_count"), ("各部門本期稽核分數", "score")):
        put(r, 3, label)
        for i, d in enumerate(depts):
            sc = scores.get(d["id"])
            put(r, 4 + i, "" if sc is None else sc[key])
        r += 1

    put(r, 3, "缺失", "FF0000")
    for i, d in enumerate(depts):
        put(r, 4 + i, d["deficiency"], "FF0000")
    r += 1

    cell_map = {(c["sheet_item_id"], c["sheet_department_id"]): c for c in detail["cells"]}
    first_item_row = r
    for it in detail["items"]:
        prefix = f"{it['display_no']}." if it["display_no"] else ""
        put(r, 3, f"{prefix}{it['name']}{it['scope_note'] or ''}")
        for i, d in enumerate(depts):
            c = cell_map.get((it["id"], d["id"]))
            if c:
                put(r, 4 + i, c["comment"], colors.get(c["result_code"]))
        r += 1
    put(first_item_row, 2, "抽檢項:")

    reviews = {rv["sheet_department_id"]: rv for rv in detail["reviews"]}
    put(r, 2, "覆核")
    put(r, 3, "上期待補正項目")
    for i, d in enumerate(depts):
        rv = reviews.get(d["id"])
        put(r, 4 + i, rv["pending_text"] if rv else "", colors.get(rv["pending_result"]) if rv else None)
    r += 1
    put(r, 3, "結果")
    for i, d in enumerate(depts):
        rv = reviews.get(d["id"])
        put(r, 4 + i, rv["result_text"] if rv else "", colors.get(rv["result_status"]) if rv else None)
    r += 1

    put(r, 2, "稽核完成率")
    put(r, 3, detail["completion_label"])
    r += 2
    if detail["remark"]:
        put(r, 1, detail["remark"])

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 32
    for i in range(len(depts)):
        ws.column_dimensions[ws.cell(row=1, column=4 + i).column_letter].width = 26

    filename = f"{detail['title']}-{detail['period']}-{detail['company_name']}.xlsx"
    return _xlsx_response(wb, filename)


@router.get("/statistics/export")
def export_statistics(
    year: int = Query(..., ge=2000, le=2999),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    data = svc.build_statistics(db, year)
    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}分數統計"
    wrap = Alignment(wrap_text=True, vertical="top")

    r = 1
    if data.get("note"):
        c = ws.cell(row=r, column=1, value=_clean(data["note"]))
        c.alignment = wrap
        r += 2

    for block in data["blocks"]:
        ws.cell(row=r, column=1, value=_clean(block["company_name"])).font = Font(bold=True)
        for i, p in enumerate(block["periods"]):
            ws.cell(row=r, column=3 + i * 2, value=_clean(p)).font = Font(bold=True)
        r += 1
        ws.cell(row=r, column=2, value="稽核時間")
        for i, lb in enumerate(block["audited_labels"]):
            ws.cell(row=r, column=3 + i * 2, value=_clean(lb))
        r += 1
        ws.cell(row=r, column=2, value="稽核項目")
        for i, majors in enumerate(block["major_items"]):
            ws.cell(row=r, column=3 + i * 2, value=_clean("\n".join(majors))).alignment = wrap
        r += 1
        for row in block["rows"]:
            ws.cell(row=r, column=2, value=_clean(row["name"]))
            for i, cell in enumerate(row["cells"]):
                ws.cell(row=r, column=3 + i * 2, value=cell["score"])
                ws.cell(row=r, column=4 + i * 2, value=_clean(cell["label"]))
            r += 1
        ws.cell(row=r, column=2, value="稽核完成率")
        for i, cell in enumerate(block["completion"]):
            ws.cell(row=r, column=3 + i * 2, value=cell["score"])
            ws.cell(row=r, column=4 + i * 2, value=_clean(cell["label"]))
        r += 3

    ws.column_dimensions["B"].width = 20
    return _xlsx_response(wb, f"財#3系統建置稽核-{year}分數統計.xlsx")


# ════════════════════════════════════════════════════════════════════════════
# 手機版輔助端點（/m/audit-check）
# ════════════════════════════════════════════════════════════════════════════
@router.get("/my-sheets")
def my_sheets(
    limit: int = Query(12, ge=1, le=60),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """
    手機首頁用：最近幾張稽核單（含公司、期別、完成率），
    比桌面版 /periods 輕量，不回傳整張矩陣。
    """
    companies = {c.id: c.name for c in db.query(Company).all()}
    periods = {p.id: p for p in db.query(AuditPeriod).all()}
    sheets = (
        db.query(AuditSheet)
        .order_by(AuditSheet.period_id.desc(), AuditSheet.company_id)
        .limit(limit)
        .all()
    )
    out = []
    for s in sheets:
        rate, label = svc.compute_completion(db, s)
        p = periods.get(s.period_id)
        out.append({
            "id": s.id,
            "period": p.period if p else "",
            "title": p.title if p else "",
            "company_id": s.company_id,
            "company_name": companies.get(s.company_id, ""),
            "audited_on": s.audited_on,
            "status": s.status,
            "completion_rate": rate,
            "completion_label": label,
        })
    out.sort(key=lambda x: (x["period"], x["company_name"]), reverse=True)
    return out
