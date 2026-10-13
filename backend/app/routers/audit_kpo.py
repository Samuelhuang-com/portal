"""
KPO 稽核（稽核檢查子模組）— FastAPI Router
Prefix: /api/v1/audit-check/kpo

中文名稱：現場營運稽核（舊名 KPO 稽核）／ 英文名稱：Audit KPO
前端路由：/audit-check/kpo/*
對應 Excel：財#1稽核系統建構（整合／2026.11執行／各部門 Sheet）
規格：docs/SPEC_audit_kpo.md

權限（沿用稽核檢查既有三個 key，不新增 key）
    audit_check_view   檢視期別與小計
    audit_check_edit   建立期別、選定稽核項目、填寫判定與備註
    audit_check_admin  部門／稽核項目主檔維護、刪除期別

⚠️ 所有端點一律掛 require_permission（CLAUDE.md §11.6）。
⚠️ 本模組非 Ragic 同步模組，sync_tool.py / RagicConnections.tsx 一律不動。
"""
from __future__ import annotations

import re
from typing import List, Optional
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Response, status
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.time import twnow
from app.dependencies import require_permission
from app.models.audit_kpo import (
    AuditKpoDept, AuditKpoItem, AuditKpoPeriod, AuditKpoPeriodItem, AuditKpoSheet, AuditKpoSheetDept,
)
from app.models.reference_data import Company, RefDepartment
from app.models.user import User
from app.schemas.audit_kpo import (
    KpoDeptCreate, KpoDeptOut, KpoDeptUpdate, KpoItemCreate, KpoItemOut,
    KpoDeptOrderUpdate, KpoItemUpdate, KpoLineOrderUpdate, KpoLineUpdate, KpoPeriodCreate, KpoPeriodDetail,
    KpoPeriodSummary, KpoPeriodUpdate, KpoSelectionUpdate, KpoSheetCreate, KpoSheetDeptsUpdate, KpoDeptHistory, KpoDashboard,
)
from app.services import audit_kpo_service as svc
from app.services.audit_check_html import sanitize_html

router = APIRouter()

VIEW = "audit_check_view"
EDIT = "audit_check_edit"
ADMIN = "audit_check_admin"


def _404(detail: str = "查無資料") -> HTTPException:
    return HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=detail)


def _409(detail: str) -> HTTPException:
    return HTTPException(status_code=status.HTTP_409_CONFLICT, detail=detail)


def _get_period(db: Session, period_id: int) -> AuditKpoPeriod:
    p = db.get(AuditKpoPeriod, period_id)
    if p is None:
        raise _404("查無此稽核期別")
    return p


# ════════════════════════════════════════════════════════════════════════════
# 部門
# ════════════════════════════════════════════════════════════════════════════
def _dept_out(db: Session, d: AuditKpoDept) -> KpoDeptOut:
    row = (
        db.query(
            func.count(AuditKpoItem.id),
            func.coalesce(func.sum(AuditKpoItem.score), 0),
        )
        .filter(AuditKpoItem.department_id == d.id, AuditKpoItem.is_active.is_(True))
        .one()
    )
    key_cnt = (
        db.query(func.count(AuditKpoItem.id))
        .filter(AuditKpoItem.department_id == d.id, AuditKpoItem.is_active.is_(True),
                AuditKpoItem.is_key.is_(True))
        .scalar()
    )
    return KpoDeptOut(
        id=d.id, name=d.name, display_name=d.display_name,
        company_id=d.company_id, company_name=d.company.name if d.company else None,
        ref_department_id=d.ref_department_id, source_title=d.source_title, sort_order=d.sort_order,
        is_active=d.is_active, item_count=int(row[0]), key_count=int(key_cnt or 0),
        total_score=int(row[1] or 0),
    )


@router.get("/departments", response_model=List[KpoDeptOut])
def list_departments(
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    svc.seed_kpo_master(db)
    depts = db.query(AuditKpoDept).order_by(AuditKpoDept.sort_order, AuditKpoDept.id).all()
    return [_dept_out(db, d) for d in depts]


def _resolve_ref_department(db: Session, company_id: int, ref_department_id: int) -> RefDepartment:
    """部門一律取自 settings/company-departments（2026-10-10 裁示，比照財#3 稽核檢查）：
    必須存在、屬於所選公司別、且未停用。"""
    rd = db.get(RefDepartment, ref_department_id)
    if rd is None:
        raise HTTPException(status_code=400, detail="查無此部門")
    if rd.company_id != company_id:
        raise HTTPException(status_code=400, detail=f"{rd.name} 不屬於所選的公司別")
    if not rd.is_active:
        raise HTTPException(status_code=400, detail=f"{rd.name} 已停用")
    return rd


def _assert_dept_unique(db: Session, company_id: Optional[int], name: str, exclude_id: Optional[int] = None) -> None:
    q = db.query(AuditKpoDept.id).filter(AuditKpoDept.name == name)
    q = q.filter(AuditKpoDept.company_id.is_(None)) if company_id is None else q.filter(AuditKpoDept.company_id == company_id)
    if exclude_id is not None:
        q = q.filter(AuditKpoDept.id != exclude_id)
    if q.first():
        raise _409(f"部門「{name}」已存在")


@router.post("/departments", response_model=KpoDeptOut, status_code=201)
def create_department(
    body: KpoDeptCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    if body.ref_department_id is not None:
        if body.company_id is None:
            raise HTTPException(status_code=400, detail="請先選擇公司別")
        rd = _resolve_ref_department(db, body.company_id, body.ref_department_id)
        name, company_id, ref_id = rd.name, body.company_id, rd.id
    else:
        # 公司別先隱藏（2026-10-11）：直接以名稱建立，不連結部門主檔
        name = (body.name or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail="請輸入部門名稱")
        company_id, ref_id = None, None
        if db.query(AuditKpoDept.id).filter(AuditKpoDept.name == name).first():
            raise _409(f"部門「{name}」已存在")      # 畫面不顯示公司別，同名會分不出來 → 不分公司一律擋
    _assert_dept_unique(db, company_id, name)
    max_order = db.query(func.coalesce(func.max(AuditKpoDept.sort_order), 0)).scalar()
    d = AuditKpoDept(
        name=name, company_id=company_id, ref_department_id=ref_id,
        source_title=body.source_title,
        sort_order=body.sort_order if body.sort_order is not None else int(max_order) + 10,
    )
    db.add(d)
    db.commit()
    db.refresh(d)
    return _dept_out(db, d)


@router.put("/departments/{dept_id}", response_model=KpoDeptOut)
def update_department(
    dept_id: int,
    body: KpoDeptUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    d = db.get(AuditKpoDept, dept_id)
    if d is None:
        raise _404("查無此部門")
    data = body.model_dump(exclude_unset=True)
    if data.get("ref_department_id") is not None:
        company_id = data.get("company_id") or d.company_id
        if company_id is None:
            raise HTTPException(status_code=400, detail="請先選擇公司別")
        rd = _resolve_ref_department(db, company_id, data["ref_department_id"])
        _assert_dept_unique(db, company_id, rd.name, exclude_id=dept_id)
        d.company_id, d.ref_department_id, d.name = company_id, rd.id, rd.name
    elif "name" in data:
        if d.ref_department_id is not None:
            raise HTTPException(status_code=400, detail="已連結部門主檔，名稱由 settings/company-departments 決定；請改選部門")
        name = data["name"].strip()
        _assert_dept_unique(db, d.company_id, name, exclude_id=dept_id)
        d.name = name
    for k in ("source_title", "sort_order", "is_active"):
        if k in data:
            setattr(d, k, data[k])
    db.commit()
    db.refresh(d)
    return _dept_out(db, d)


# ════════════════════════════════════════════════════════════════════════════
# 稽核項目主檔
# ════════════════════════════════════════════════════════════════════════════
def _item_out(it: AuditKpoItem, in_use: bool) -> KpoItemOut:
    return KpoItemOut(
        id=it.id, department_id=it.department_id, item_no=it.item_no, is_key=it.is_key,
        name=it.name, score=int(it.score or 0), method=it.method,
        sort_order=it.sort_order, is_active=it.is_active, in_use=in_use,
    )


@router.get("/items", response_model=List[KpoItemOut])
def list_items(
    department_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    svc.seed_kpo_master(db)
    q = db.query(AuditKpoItem)
    if department_id is not None:
        q = q.filter(AuditKpoItem.department_id == department_id)
    items = q.order_by(AuditKpoItem.department_id, AuditKpoItem.sort_order, AuditKpoItem.id).all()
    used = {
        r[0] for r in db.query(AuditKpoPeriodItem.source_item_id)
        .filter(AuditKpoPeriodItem.source_item_id.isnot(None)).distinct().all()
    }
    return [_item_out(it, it.id in used) for it in items]


@router.post("/items", response_model=KpoItemOut, status_code=201)
def create_item(
    body: KpoItemCreate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    if db.get(AuditKpoDept, body.department_id) is None:
        raise _404("查無此部門")
    max_order = db.query(func.coalesce(func.max(AuditKpoItem.sort_order), 0)).filter(
        AuditKpoItem.department_id == body.department_id).scalar()
    max_no = db.query(func.coalesce(func.max(AuditKpoItem.item_no), 0)).filter(
        AuditKpoItem.department_id == body.department_id).scalar()
    it = AuditKpoItem(
        department_id=body.department_id,
        item_no=body.item_no if body.item_no is not None else int(max_no) + 1,
        is_key=body.is_key, name=body.name.strip(),
        score=body.score if body.score is not None else svc.default_score(body.is_key),
        method=sanitize_html(body.method) or None, sort_order=int(max_order) + 10,
    )
    db.add(it)
    db.commit()
    db.refresh(it)
    return _item_out(it, False)


@router.put("/items/{item_id}", response_model=KpoItemOut)
def update_item(
    item_id: int,
    body: KpoItemUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    it = db.get(AuditKpoItem, item_id)
    if it is None:
        raise _404("查無此稽核項目")
    data = body.model_dump(exclude_unset=True)
    if "name" in data:
        data["name"] = data["name"].strip()
    if "method" in data:
        data["method"] = sanitize_html(data["method"]) or None   # 備註為 Web 編輯器 HTML，白名單清理
    for k, v in data.items():
        setattr(it, k, v)
    db.commit()
    db.refresh(it)
    in_use = db.query(AuditKpoPeriodItem.id).filter(
        AuditKpoPeriodItem.source_item_id == it.id).first() is not None
    return _item_out(it, in_use)


@router.delete("/items/{item_id}", status_code=204)
def delete_item(
    item_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    """刪除主檔項目。已選定於歷史期別的快照保留（source_item_id 轉 NULL），不受影響。"""
    it = db.get(AuditKpoItem, item_id)
    if it is None:
        raise _404("查無此稽核項目")
    db.query(AuditKpoPeriodItem).filter(
        AuditKpoPeriodItem.source_item_id == item_id
    ).update({AuditKpoPeriodItem.source_item_id: None}, synchronize_session=False)
    db.delete(it)
    db.commit()


# ════════════════════════════════════════════════════════════════════════════
# 期別
# ════════════════════════════════════════════════════════════════════════════
@router.get("/dashboard", response_model=KpoDashboard)
def get_dashboard(
    year: Optional[int] = Query(None, ge=2000, le=2100),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """年度 Dashboard：公司 × 部門 × 月份的本期分數（同月多期取平均、列全部啟用部門）。"""
    return svc.dashboard(db, year)


@router.get("/statistics")
def get_statistics(
    year: Optional[int] = Query(None, ge=2000, le=2100),
    company_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """分數統計（2026-10-13）：公司 × 部門 × 期別的本期分數，附「N項符合 / M項選定」與查核完成列。"""
    return svc.statistics(db, year, company_id)


@router.get("/statistics/flagged")
def get_flagged(
    year: Optional[int] = Query(None, ge=2000, le=2100),
    include_pending: bool = Query(False, description="True 連「未查核」項目一併列出"),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """缺失與建議清單：不符合項（含查核備註）。"""
    return svc.flagged(db, year, include_pending)


@router.get("/statistics/export")
def export_statistics(
    year: Optional[int] = Query(None, ge=2000, le=2100),
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """匯出分數統計 Excel：每家公司一個 Sheet；每期兩欄（本期分數／符合÷選定），最後一列為查核完成。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from app.routers.audit_check import _clean, _xlsx_response

    data = svc.statistics(db, year)
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    used: set = set()
    for blk in data["blocks"]:
        title = _clean(blk["company_name"])[:28] or "公司"
        base, n = title, 2
        while title in used:
            title, n = f"{base}{n}", n + 1
        used.add(title)
        ws = wb.create_sheet(title)
        ws.cell(row=1, column=1, value=f"{blk['company_name']}　{data['year']} 年　現場營運稽核分數統計").font = bold
        ws.cell(row=2, column=1, value="部門").font = bold
        for i, lb in enumerate(blk["periods"]):
            ws.cell(row=2, column=2 + i * 2, value=_clean(lb)).font = bold
            ws.cell(row=2, column=3 + i * 2, value="符合/選定").font = bold
        r = 3
        for row in blk["rows"]:
            ws.cell(row=r, column=1, value=_clean(row["name"]))
            for i, c in enumerate(row["cells"]):
                ws.cell(row=r, column=2 + i * 2, value=c["score"])
                ws.cell(row=r, column=3 + i * 2, value=f"{c['pass_count']}/{c['item_count']}" if c["item_count"] else "本月無")
            r += 1
        ws.cell(row=r, column=1, value="查核完成（已查核/選定）").font = bold
        for i, c in enumerate(blk["completion"]):
            ws.cell(row=r, column=2 + i * 2, value=c["score"])
            ws.cell(row=r, column=3 + i * 2, value=_clean(c["label"]))
        ws.column_dimensions["A"].width = 24
    if not data["blocks"]:
        ws = wb.create_sheet("無資料")
        ws.cell(row=1, column=1, value=f"{data['year']} 年尚無現場營運稽核資料").alignment = Alignment(wrap_text=True)
    return _xlsx_response(wb, f"現場營運稽核-{data['year']}分數統計.xlsx")


@router.get("/periods", response_model=List[KpoPeriodSummary])
def list_periods(
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    periods = db.query(AuditKpoPeriod).order_by(
        AuditKpoPeriod.period.desc(), AuditKpoPeriod.id.desc()).all()
    return [svc.build_period_summary(p) for p in periods]


@router.post("/periods", response_model=KpoPeriodDetail, status_code=201)
def create_period(
    body: KpoPeriodCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    p = AuditKpoPeriod(
        period=body.period, title=body.title.strip() or "財#1稽核系統建構",
        note=body.note, created_by=str(user.id),
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.get("/periods/{period_id}", response_model=KpoPeriodDetail)
def get_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    return svc.build_period_detail(db, _get_period(db, period_id))


@router.put("/periods/{period_id}", response_model=KpoPeriodDetail)
def update_period(
    period_id: int,
    body: KpoPeriodUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    p = _get_period(db, period_id)
    for k, v in body.model_dump(exclude_unset=True).items():
        if k in ("period", "title") and not v:
            continue
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.delete("/periods/{period_id}", status_code=204)
def delete_period(
    period_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    p = _get_period(db, period_id)
    db.delete(p)
    db.commit()


def _get_sheet(p: AuditKpoPeriod, sheet_id: int) -> AuditKpoSheet:
    for sh in p.sheets:
        if sh.id == sheet_id:
            return sh
    raise _404("查無此公司稽核單")


def _check_master_depts(db: Session, department_ids: List[int]) -> List[int]:
    """受稽部門＝現場營運稽核項目的部門（audit_kpo_departments）；不依公司別過濾。"""
    ids = list(dict.fromkeys(department_ids))
    rows = {d.id: d for d in db.query(AuditKpoDept).filter(AuditKpoDept.id.in_(ids or [0])).all()}
    bad = [i for i in ids if i not in rows]
    if bad:
        raise HTTPException(status_code=400, detail=f"查無部門：{bad}")
    stopped = [rows[i].display_name for i in ids if not rows[i].is_active]
    if stopped:
        raise HTTPException(status_code=400, detail=f"部門已停用：{'、'.join(stopped)}")
    return ids


@router.post("/periods/{period_id}/sheets", response_model=KpoPeriodDetail, status_code=201)
def create_sheet(
    period_id: int,
    body: KpoSheetCreate,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(EDIT)),
):
    """新增公司別稽核單（一期一家公司一張）＋選受稽部門（可跨公司任選主檔部門）。"""
    p = _get_period(db, period_id)
    company = db.get(Company, body.company_id)
    if company is None:
        raise HTTPException(status_code=400, detail="查無此公司別")
    if any(sh.company_id == body.company_id for sh in p.sheets):
        raise _409(f"本期已建立「{company.name}」的稽核單")
    ids = _check_master_depts(db, body.department_ids)
    sh = AuditKpoSheet(period_id=p.id, company_id=body.company_id, created_by=str(user.id))
    for idx, dept_id in enumerate(ids):
        sh.departments.append(AuditKpoSheetDept(department_id=dept_id, sort_order=(idx + 1) * 10))
    p.sheets.append(sh)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise _409("稽核單重複，請重新整理後再試")
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.delete("/periods/{period_id}/sheets/{sheet_id}", response_model=KpoPeriodDetail)
def delete_sheet(
    period_id: int,
    sheet_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(ADMIN)),
):
    """刪除公司稽核單（其受稽部門、選定項目與判定一併刪除）。"""
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    p.sheets.remove(sh)
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.put("/periods/{period_id}/sheets/{sheet_id}/departments", response_model=KpoPeriodDetail)
def update_sheet_departments(
    period_id: int,
    sheet_id: int,
    body: KpoSheetDeptsUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    """整組替換受稽部門：新增的接在最後；移除的連同選定項目一併刪除；保留的維持原順序。"""
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    ids = _check_master_depts(db, body.department_ids)
    have = {sd.department_id: sd for sd in sh.departments}
    for dept_id, sd in list(have.items()):
        if dept_id not in ids:
            for ln in [x for x in sh.lines if x.department_id == dept_id]:
                sh.lines.remove(ln)
            sh.departments.remove(sd)
    next_order = max([sd.sort_order or 0 for sd in have.values()], default=0)
    for dept_id in ids:
        if dept_id not in have:
            next_order += 10
            sh.departments.append(AuditKpoSheetDept(department_id=dept_id, sort_order=next_order))
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.get("/periods/{period_id}/sheets/{sheet_id}/departments/{department_id}/history",
            response_model=KpoDeptHistory)
def get_dept_history(
    period_id: int,
    sheet_id: int,
    department_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(VIEW)),
):
    """「選擇項目」視窗用：同公司、同部門在其他期別的期別數／分數／平均分數，與各項目曾選過的期別＋判定。"""
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    return svc.dept_history(db, p, sh, department_id)


@router.put("/periods/{period_id}/sheets/{sheet_id}/selection", response_model=KpoPeriodDetail)
def update_selection(
    period_id: int,
    sheet_id: int,
    body: KpoSelectionUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    """
    整組替換某稽核單、某受稽部門選定的項目：
      - 新勾選的 → 由主檔複製一份快照（含 ★、分數、備註）
      - 取消勾選的 → 移除該列（其判定與備註一併消失）
      - 原本就選定的 → 保留不動（本期已編輯的內容與判定不會被主檔覆蓋）
    """
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    if not any(sd.department_id == body.department_id for sd in sh.departments):
        raise HTTPException(status_code=400, detail="此部門不是該稽核單的受稽部門")
    wanted = list(dict.fromkeys(body.item_ids))
    masters = {
        it.id: it for it in db.query(AuditKpoItem).filter(
            AuditKpoItem.department_id == body.department_id,
            AuditKpoItem.id.in_(wanted or [0]),
        ).all()
    }
    bad = [i for i in wanted if i not in masters]
    if bad:
        raise HTTPException(status_code=400, detail=f"稽核項目不屬於該部門或不存在：{bad}")

    mine = [ln for ln in sh.lines if ln.department_id == body.department_id]
    next_order = max([ln.sort_order or 0 for ln in mine], default=0)
    existing = {ln.source_item_id: ln for ln in mine if ln.source_item_id is not None}
    for sid, ln in list(existing.items()):
        if sid not in masters:
            sh.lines.remove(ln)
    # 新勾選的項目接在該部門現有順序之後（依主檔順序），不會插進使用者調整過的順序中間
    for sid in sorted((i for i in wanted if i not in existing), key=lambda i: (masters[i].sort_order, i)):
        m = masters[sid]
        next_order += 10
        sh.lines.append(AuditKpoPeriodItem(
            period_id=p.id, department_id=body.department_id, source_item_id=m.id,
            item_no=m.item_no, is_key=m.is_key, name=m.name, score=m.score,
            method=m.method, sort_order=next_order,
        ))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise _409("選定項目重複，請重新整理後再試")
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.put("/periods/{period_id}/sheets/{sheet_id}/dept-order", response_model=KpoPeriodDetail)
def update_dept_order(
    period_id: int,
    sheet_id: int,
    body: KpoDeptOrderUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    """儲存該稽核單的受稽部門顯示順序（只影響這張稽核單）。"""
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    mine = {sd.department_id: sd for sd in sh.departments}
    ids = list(dict.fromkeys(body.department_ids))
    bad = [i for i in ids if i not in mine]
    if bad:
        raise HTTPException(status_code=400, detail=f"部門不屬於該稽核單：{bad}")
    rest = sorted((sd for i, sd in mine.items() if i not in ids), key=lambda x: (x.sort_order or 0, x.id))
    for idx, sd in enumerate([mine[i] for i in ids] + rest):
        sd.sort_order = (idx + 1) * 10
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.put("/periods/{period_id}/sheets/{sheet_id}/lines-order", response_model=KpoPeriodDetail)
def update_lines_order(
    period_id: int,
    sheet_id: int,
    body: KpoLineOrderUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    """儲存某稽核單、某部門的項目順序（只影響這張稽核單的這個部門）。"""
    p = _get_period(db, period_id)
    sh = _get_sheet(p, sheet_id)
    mine = {ln.id: ln for ln in sh.lines if ln.department_id == body.department_id}
    ids = list(dict.fromkeys(body.line_ids))
    bad = [i for i in ids if i not in mine]
    if bad:
        raise HTTPException(status_code=400, detail=f"選定項目不屬於該稽核單的此部門：{bad}")
    rest = sorted((ln for i, ln in mine.items() if i not in ids),
                  key=lambda x: (x.sort_order or 0, x.id))
    for idx, ln in enumerate([mine[i] for i in ids] + rest):
        ln.sort_order = (idx + 1) * 10
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.put("/periods/{period_id}/lines/{line_id}", response_model=KpoPeriodDetail)
def update_line(
    period_id: int,
    line_id: int,
    body: KpoLineUpdate,
    db: Session = Depends(get_db),
    _: User = Depends(require_permission(EDIT)),
):
    """編輯本期選定項目的內容／判定／備註（只影響本期這一列，不動主檔）。"""
    p = _get_period(db, period_id)
    ln = db.get(AuditKpoPeriodItem, line_id)
    if ln is None or ln.period_id != p.id:
        raise _404("查無此選定項目")
    data = body.model_dump(exclude_unset=True)
    if "result" in data:
        ln.result = data.pop("result") or None
    if "name" in data:
        data["name"] = data["name"].strip()
    if "method" in data:
        data["method"] = sanitize_html(data["method"]) or None   # 備註為 Web 編輯器 HTML，白名單清理
    if "remark" in data:
        data["remark"] = sanitize_html(data["remark"]) or None   # 查核備註同為 Web 編輯器 HTML
    for k, v in data.items():
        setattr(ln, k, v)
    db.commit()
    db.refresh(p)
    return svc.build_period_detail(db, p)


@router.get("/periods/{period_id}/sheets/{sheet_id}/departments/{department_id}/inspection-docx")
def export_inspection_docx(
    period_id: int,
    sheet_id: int,
    department_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_permission(VIEW)),
):
    """單一受稽單位「內部稽核報告」Word（2026-10-12）——版型同稽核檢查 sheets 的匯出 Word。
    填表日期＝今天；稽核人員＝匯出的使用者。"""
    try:
        from app.services.audit_kpo_docx import build_kpo_inspection_docx
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="伺服器尚未安裝 python-docx，請在 backend 執行 pip install -r requirements.txt 後重啟",
        )
    p = _get_period(db, period_id)
    _get_sheet(p, sheet_id)
    detail = svc.build_period_detail(db, p)
    sh = next(s for s in detail["sheets"] if s["id"] == sheet_id)
    idx = next((i for i, b in enumerate(sh["departments"], 1) if b["department_id"] == department_id), None)
    if idx is None:
        raise _404("此公司別稽核單沒有這個受稽單位")
    block = sh["departments"][idx - 1]
    data = build_kpo_inspection_docx(
        detail, sh["company_name"], block, idx, user.full_name or "", twnow().date().strftime("%Y.%m.%d"),
    )
    filename_cn = f"內部稽核報告-{detail['period']}-{sh['company_name']}-{block['department_name']}.docx"
    ascii_name = re.sub(r"[^A-Za-z0-9._-]+", "_", filename_cn).strip("_") or "export.docx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8\'\'{quote(filename_cn, safe='')}"},
    )
