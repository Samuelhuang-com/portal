# -*- coding: utf-8 -*-
"""
SQL 查詢工具 service（settings/sql-console，僅系統管理員）— PostgreSQL 專用

安全層次（由外而內）：
  1. app/core/sql_guard.check_sql：只放行 SELECT / WITH / EXPLAIN 與「有 WHERE 的單表 DELETE」
  2. SELECT：在 `SET TRANSACTION READ ONLY` 交易內執行（資料庫本身拒絕任何寫入）＋ statement_timeout ＋ 筆數上限
  3. DELETE：兩段式
       ① 預覽：在交易內「真的執行」DELETE … RETURNING * 取得筆數與被刪資料 → ROLLBACK
          （超過 MAX_DELETE_ROWS 直接拒絕）→ 回傳簽章 token（綁定使用者＋SQL＋預覽筆數，5 分鐘內有效）
       ② 確認：帶 token 再執行一次，筆數必須與預覽相同才 COMMIT；被刪資料完整備份寫入 audit_logs.extra
          （刪除與稽核紀錄在同一個交易，要嘛一起成功要嘛一起回復）
  4. 所有查詢／匯出／刪除／被擋的嘗試都寫入 audit_logs（action = sql_console.*）

注意：用 psycopg 原生 cursor 執行（不經過 SQLAlchemy text()），所以 SQL 內的 `:` 與 `%` 都不會被當成參數。
"""
from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import math
import re
import time
import uuid
from datetime import date, datetime, time as dtime, timedelta
from decimal import Decimal
from typing import Any, Optional

from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.excel_utils import xlsx_safe
from app.core.sql_guard import GuardResult, SqlGuardError, check_sql
from app.core.time import twnow
from app.models.audit_log import AuditLog
from app.models.user import User

VIEW_DEFAULT_ROWS = 1000
VIEW_MAX_ROWS = 5000
EXPORT_MAX_ROWS = 50_000
SELECT_TIMEOUT_MS = 30_000
EXPORT_TIMEOUT_MS = 120_000
DELETE_TIMEOUT_MS = 60_000
MAX_DELETE_ROWS = 1000          # 單次 DELETE 最多影響筆數
DELETE_SAMPLE_ROWS = 20         # 預覽時回給畫面看的筆數
TOKEN_TTL_SEC = 300
LOG_SQL_MAX = 4000

# 結果集中這類欄位一律遮成 ***（畫面與匯出；備份不遮）
_MASK_COL_RE = re.compile(r"password|passwd|secret|token|api_?key|hash|credential", re.I)

ACTION_SELECT = "sql_console.select"
ACTION_EXPORT = "sql_console.export"
ACTION_DELETE = "sql_console.delete"
ACTION_BLOCKED = "sql_console.blocked"


class SqlConsoleError(Exception):
    """可直接顯示給使用者的執行錯誤（400）。"""


# ── 共用 ─────────────────────────────────────────────────────────────────────

def actor_of(user: User) -> dict:
    return {"id": user.id, "tenant_id": getattr(user, "tenant_id", None), "email": user.email}


def _ensure_pg(db: Session) -> None:
    if db.get_bind().dialect.name != "postgresql":
        raise SqlConsoleError("SQL 查詢工具只支援 PostgreSQL（目前連到的不是 PostgreSQL，請檢查 .env 的 DATABASE_URL）")


def _cursor(db: Session):
    """取得與 Session 同一個交易的 psycopg 原生 cursor。"""
    return db.connection().connection.driver_connection.cursor()


def _pg_msg(e: Exception) -> str:
    diag = getattr(e, "diag", None)
    msg = getattr(diag, "message_primary", None) or str(e)
    detail = getattr(diag, "message_detail", None)
    hint = getattr(diag, "message_hint", None)
    out = msg.strip()
    if detail:
        out += f"\n{detail.strip()}"
    if hint:
        out += f"\n提示：{hint.strip()}"
    return out


def _jsonable(v: Any) -> Any:
    """轉成可 JSON 序列化、且人看得懂的值。"""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else str(v)
    if isinstance(v, Decimal):
        return float(v) if v.is_finite() else str(v)
    if isinstance(v, (datetime, date, dtime)):
        return v.isoformat(sep=" ") if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, timedelta):
        return str(v)
    if isinstance(v, uuid.UUID):
        return str(v)
    if isinstance(v, (bytes, bytearray, memoryview)):
        return f"<binary {len(bytes(v))} bytes>"
    if isinstance(v, (dict, list, tuple)):
        return json.dumps(v, ensure_ascii=False, default=_jsonable)
    return str(v)


def _mask_rows(columns: list[str], rows: list[list[Any]]) -> list[list[Any]]:
    idx = [i for i, c in enumerate(columns) if _MASK_COL_RE.search(c or "")]
    if not idx:
        return rows
    for r in rows:
        for i in idx:
            if r[i] is not None:
                r[i] = "***"
    return rows


def _write_log(db: Session, actor: dict, action: str, ip: Optional[str],
               resource_id: Optional[str] = None, extra: Optional[dict] = None) -> str:
    log = AuditLog(
        user_id=actor["id"], tenant_id=actor.get("tenant_id"),
        action=action, resource_type="sql_console", resource_id=(resource_id or None),
        ip_address=ip, extra=extra or {},
    )
    db.add(log)
    db.commit()
    return log.id


def log_blocked(db: Session, actor: dict, sql: str, reason: str, ip: Optional[str]) -> None:
    try:
        _write_log(db, actor, ACTION_BLOCKED, ip, extra={"sql": (sql or "")[:LOG_SQL_MAX], "reason": reason})
    except Exception:
        db.rollback()


# ── SELECT ───────────────────────────────────────────────────────────────────

def _run_readonly(db: Session, sql: str, max_rows: int, timeout_ms: int) -> tuple[list[str], list[list[Any]], bool, int]:
    """回傳 (欄位, 資料列, 是否被截斷, 耗時 ms)。資料列已轉成 JSON 安全的值。"""
    _ensure_pg(db)
    db.rollback()                       # 結束 dependency 先前開的交易，確保 READ ONLY 是新交易的第一句
    cur = None
    try:
        cur = _cursor(db)
        cur.execute("SET TRANSACTION READ ONLY")
        cur.execute(f"SET LOCAL statement_timeout = {int(timeout_ms)}")
        t0 = time.perf_counter()
        cur.execute(sql)
        if cur.description is None:
            raise SqlConsoleError("此敘述沒有回傳結果集")
        columns = [d.name for d in cur.description]
        fetched = cur.fetchmany(max_rows + 1)
        elapsed = int((time.perf_counter() - t0) * 1000)
    except SqlConsoleError:
        raise
    except Exception as e:  # noqa: BLE001 — 轉成使用者可讀訊息
        raise SqlConsoleError(_pg_msg(e)) from e
    finally:
        try:
            if cur is not None:
                cur.close()
        finally:
            db.rollback()
    truncated = len(fetched) > max_rows
    rows = [[_jsonable(v) for v in r] for r in fetched[:max_rows]]
    return columns, _mask_rows(columns, rows), truncated, elapsed


def run_select(db: Session, actor: dict, gr: GuardResult, limit: int, ip: Optional[str]) -> dict:
    limit = max(1, min(int(limit or VIEW_DEFAULT_ROWS), VIEW_MAX_ROWS))
    columns, rows, truncated, elapsed = _run_readonly(db, gr.sql, limit, SELECT_TIMEOUT_MS)
    _write_log(db, actor, ACTION_SELECT, ip, extra={
        "sql": gr.sql[:LOG_SQL_MAX], "row_count": len(rows), "truncated": truncated, "elapsed_ms": elapsed,
    })
    return {"kind": "select", "columns": columns, "rows": rows, "row_count": len(rows),
            "truncated": truncated, "limit": limit, "elapsed_ms": elapsed}


# ── DELETE（兩段式）─────────────────────────────────────────────────────────

def _token_key() -> bytes:
    return (settings.JWT_SECRET_KEY or settings.SECRET_KEY).encode("utf-8")


def _sign(uid: str, sql: str, affected: int, exp: int) -> str:
    msg = f"{uid}|{hashlib.sha256(sql.encode('utf-8')).hexdigest()}|{affected}|{exp}"
    sig = hmac.new(_token_key(), msg.encode("utf-8"), hashlib.sha256).hexdigest()
    return f"{affected}.{exp}.{sig}"


def _verify_token(uid: str, sql: str, token: str) -> int:
    """驗證確認用 token，回傳預覽時的筆數。"""
    try:
        affected_s, exp_s, sig = token.split(".", 2)
        affected, exp = int(affected_s), int(exp_s)
    except Exception:
        raise SqlConsoleError("確認碼格式不正確，請重新預覽") from None
    if time.time() > exp:
        raise SqlConsoleError("確認碼已過期（超過 5 分鐘），請重新預覽")
    expected = _sign(uid, sql, affected, exp)
    if not hmac.compare_digest(expected, token):
        raise SqlConsoleError("確認碼與目前的 SQL 不符（SQL 被修改過或不是同一位使用者），請重新預覽")
    return affected


def _run_delete_in_tx(db: Session, gr: GuardResult) -> tuple[int, list[str], list[list[Any]]]:
    """在目前交易內執行 DELETE … RETURNING *，不 commit。回傳 (筆數, 欄位, 被刪的原始資料列)。"""
    _ensure_pg(db)
    cur = None
    try:
        cur = _cursor(db)
        cur.execute(f"SET LOCAL statement_timeout = {DELETE_TIMEOUT_MS}")
        # 換行：避免 SQL 結尾的 -- 註解吃掉 RETURNING
        cur.execute(gr.sql + "\nRETURNING *")
        if cur.description is None:
            raise SqlConsoleError("DELETE 沒有回傳被刪資料，已取消")
        columns = [d.name for d in cur.description]
        rows = cur.fetchall()
        affected = cur.rowcount
    except SqlConsoleError:
        db.rollback()
        raise
    except Exception as e:  # noqa: BLE001
        db.rollback()
        raise SqlConsoleError(_pg_msg(e)) from e
    finally:
        if cur is not None:
            try:
                cur.close()
            except Exception:
                pass
    if affected > MAX_DELETE_ROWS:
        db.rollback()
        raise SqlConsoleError(
            f"此 DELETE 會影響 {affected} 筆，超過單次上限 {MAX_DELETE_ROWS} 筆，已全部復原。請加上更精確的 WHERE 條件。"
        )
    return affected, columns, rows


def preview_delete(db: Session, actor: dict, gr: GuardResult) -> dict:
    db.rollback()
    affected, columns, rows = _run_delete_in_tx(db, gr)
    db.rollback()                       # 預覽：一律復原
    sample = _mask_rows(columns, [[_jsonable(v) for v in r] for r in rows[:DELETE_SAMPLE_ROWS]])
    exp = int(time.time()) + TOKEN_TTL_SEC
    return {
        "kind": "delete_preview", "table": gr.table, "affected": affected,
        "columns": columns, "sample_rows": sample, "sample_size": len(sample),
        "token": _sign(actor["id"], gr.sql, affected, exp), "expires_in": TOKEN_TTL_SEC,
        "max_delete_rows": MAX_DELETE_ROWS,
    }


def commit_delete(db: Session, actor: dict, gr: GuardResult, token: str, ip: Optional[str]) -> dict:
    expected = _verify_token(actor["id"], gr.sql, token)
    db.rollback()
    affected, columns, rows = _run_delete_in_tx(db, gr)
    if affected != expected:
        db.rollback()
        raise SqlConsoleError(f"預覽時是 {expected} 筆，現在變成 {affected} 筆（資料在這段時間被異動），已全部復原。請重新預覽。")
    backup = [{c: _jsonable(v) for c, v in zip(columns, r)} for r in rows]
    try:
        log = AuditLog(
            user_id=actor["id"], tenant_id=actor.get("tenant_id"), action=ACTION_DELETE,
            resource_type="sql_console", resource_id=gr.table, ip_address=ip,
            extra={"sql": gr.sql[:LOG_SQL_MAX], "table": gr.table, "affected": affected,
                   "columns": columns, "deleted_rows": backup},
        )
        db.add(log)
        db.commit()                     # 刪除 ＋ 備份紀錄同一個交易
    except Exception as e:  # noqa: BLE001
        db.rollback()
        raise SqlConsoleError(f"寫入稽核備份失敗，刪除已復原：{e}") from e
    return {"kind": "delete_done", "table": gr.table, "affected": affected, "log_id": log.id}


# ── 匯出 ─────────────────────────────────────────────────────────────────────

def _excel_value(v: Any) -> Any:
    if isinstance(v, str):
        return xlsx_safe(v)
    return v


def export_result(db: Session, actor: dict, gr: GuardResult, fmt: str, ip: Optional[str]) -> dict:
    fmt = (fmt or "").lower()
    if fmt not in ("xlsx", "csv"):
        raise SqlConsoleError("匯出格式只支援 xlsx 或 csv")
    columns, rows, truncated, elapsed = _run_readonly(db, gr.sql, EXPORT_MAX_ROWS, EXPORT_TIMEOUT_MS)
    stamp = twnow().strftime("%Y%m%d_%H%M%S")
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(columns)
        for r in rows:
            w.writerow(["" if v is None else v for v in r])
        data = buf.getvalue().encode("utf-8-sig")        # BOM：Excel 直接開啟中文才不會亂碼
        ctype = "text/csv; charset=utf-8"
    else:
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("result")
        ws.append([_excel_value(c) for c in columns])
        for r in rows:
            ws.append([_excel_value(v) for v in r])
        out = io.BytesIO()
        wb.save(out)
        data = out.getvalue()
        ctype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    _write_log(db, actor, ACTION_EXPORT, ip, extra={
        "sql": gr.sql[:LOG_SQL_MAX], "format": fmt, "row_count": len(rows), "truncated": truncated, "elapsed_ms": elapsed,
    })
    return {"data": data, "content_type": ctype, "filename": f"sql_result_{stamp}.{fmt}",
            "row_count": len(rows), "truncated": truncated}


# ── 資料表結構／稽核紀錄 ─────────────────────────────────────────────────────

def list_tables(db: Session) -> list[dict]:
    _ensure_pg(db)
    cols = db.execute(text(
        "SELECT c.table_name, c.column_name, c.data_type "
        "FROM information_schema.columns c "
        "JOIN information_schema.tables t ON t.table_schema = c.table_schema "
        "  AND t.table_name = c.table_name AND t.table_type = 'BASE TABLE' "
        "WHERE c.table_schema = 'public' ORDER BY c.table_name, c.ordinal_position"
    )).all()
    est = dict(db.execute(text(
        "SELECT relname, GREATEST(reltuples, 0)::bigint FROM pg_class "
        "WHERE relkind = 'r' AND relnamespace = 'public'::regnamespace"
    )).all())
    out: dict[str, dict] = {}
    for t, c, dt in cols:
        out.setdefault(t, {"table": t, "est_rows": int(est.get(t, 0)), "columns": []})["columns"].append({"name": c, "type": dt})
    return list(out.values())


def list_logs(db: Session, limit: int = 100) -> list[dict]:
    limit = max(1, min(int(limit or 100), 500))
    rows = (
        db.query(AuditLog, User.email)
        .outerjoin(User, User.id == AuditLog.user_id)
        .filter(AuditLog.action.like("sql_console.%"))
        .order_by(AuditLog.created_at.desc())
        .limit(limit)
        .all()
    )
    result = []
    for log, email in rows:
        ex = log.extra or {}
        result.append({
            "id": log.id,
            "created_at": log.created_at.isoformat(sep=" ", timespec="seconds") if log.created_at else None,
            "email": email,
            "action": log.action.replace("sql_console.", ""),
            "table": ex.get("table"),
            "affected": ex.get("affected"),
            "row_count": ex.get("row_count"),
            "reason": ex.get("reason"),
            "sql": ex.get("sql"),
            "has_backup": bool(ex.get("deleted_rows")),
            "ip": log.ip_address,
        })
    return result


def get_log_detail(db: Session, log_id: str) -> Optional[dict]:
    log = db.query(AuditLog).filter(AuditLog.id == log_id, AuditLog.action.like("sql_console.%")).first()
    if not log:
        return None
    ex = log.extra or {}
    cols = ex.get("columns") or []
    rows = [[r.get(c) for c in cols] for r in (ex.get("deleted_rows") or [])]
    return {"id": log.id, "action": log.action.replace("sql_console.", ""), "sql": ex.get("sql"),
            "table": ex.get("table"), "affected": ex.get("affected"), "columns": cols, "rows": rows}


def guard(sql: str) -> GuardResult:
    """薄包裝，讓 router 只需 import 一個模組。"""
    return check_sql(sql)


__all__ = [
    "SqlConsoleError", "SqlGuardError", "actor_of", "guard", "log_blocked", "run_select",
    "preview_delete", "commit_delete", "export_result", "list_tables", "list_logs", "get_log_detail",
    "VIEW_DEFAULT_ROWS", "VIEW_MAX_ROWS", "EXPORT_MAX_ROWS", "MAX_DELETE_ROWS",
]
