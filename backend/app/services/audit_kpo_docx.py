"""
KPO 現場營運稽核 — 單一受稽單位「內部稽核報告」Word 匯出（2026-10-12）

版型、字型、字級與稽核檢查 sheets 的「匯出 Word」完全相同（共用 audit_check_docx 的工具與常數）：
  logo／標題 16pt／填表日期 14pt／其餘 11pt（微軟正黑體）／「本次稽核結果」與統計表 14pt／簽核列三格。

欄位對應（2026-10-12 使用者裁示）：
  項次＝項數（重點項前綴 ★）｜內部稽核項目＝品質項目｜稽核紀錄＝查核備註（富文字）
  ｜稽核結果＝符合 V／不符合 X（未判定留白）｜最右欄＝備註（查核方式）（富文字）
列入範圍：該單位本期所有選定項目（不分大項，一層列出）。
缺失＝自動列出不符合項的「項數 品質項目」；統計＝選定項數／不符合項數／本期分數（score100，N分）。
數字一律取 audit_kpo_service 的系統計算結果（subtotal），不在此重算。
"""
from __future__ import annotations

import io
from typing import Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from app.services import audit_check_docx as base
from app.services.audit_check_docx import (
    BODY_PT, COLS, DATE_PT, FONT, LOGO_PATH, RED, RESULT_PT, TITLE_PT,
    _fix_widths, _font, _highlight, _merge, _row_height, _score_color, _shade, _write, _write_html,
)
from docx.oxml.ns import qn


def _item_no(line: dict) -> str:
    no = line.get("item_no")
    s = "" if no is None else str(no)
    return f"★{s}" if line.get("is_key") else s


def build_kpo_inspection_docx(
    period: dict,
    company_name: str,
    block: dict,
    dept_index: int,
    auditor_name: str,
    fill_date: str,
) -> bytes:
    """period：build_period_detail() 結果；block：該單位的 KPO 部門區塊（含 lines／subtotal）。"""
    dept_name = block["department_name"]
    lines = block["lines"]
    sub = block["subtotal"]
    period_s = period.get("period") or ""

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.0)
    sec.bottom_margin = Cm(0.75)
    sec.header_distance = Cm(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)

    hp = sec.header.paragraphs[0]
    if LOGO_PATH.exists():
        hp.add_run().add_picture(str(LOGO_PATH), width=Cm(4.2))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("內部稽核報告"), TITLE_PT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(p.add_run(f"填表日期：{fill_date}"), DATE_PT)

    t = doc.add_table(rows=0, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    def add_row(height: float = 0.75):
        r = t.add_row()
        _row_height(r, height)
        return r

    r = add_row(0.9)
    _write(_merge(r, 0, 2), f"受稽單位：{company_name}", BODY_PT)
    _write(_merge(r, 3, 4), f"文件編號：{period_s.replace('-', '')}-{dept_index}", BODY_PT)
    r = add_row()
    _write(_merge(r, 0, 2), f"受稽人員：{dept_name}", BODY_PT)
    _write(_merge(r, 3, 4), f"稽核人員：{auditor_name}", BODY_PT)
    r = add_row()
    _write(_merge(r, 0, 4), f"文件名稱：{company_name}現場營運稽核-{dept_name}", BODY_PT)

    r = add_row(1.0)
    for idx, label in enumerate(["項次", "內部稽核項目", "稽核紀錄", "稽核結果", "備註（查核方式）"]):
        _write(r.cells[idx], label, BODY_PT, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(r.cells[idx])

    for ln in lines:
        rr = add_row(0.9)
        _write(rr.cells[0], _item_no(ln), BODY_PT, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(rr.cells[1], ln.get("name") or "", BODY_PT)
        _write_html(rr.cells[2], ln.get("remark") or "", BODY_PT)
        res = ln.get("result")
        _write(rr.cells[3], "V" if res == "pass" else "X" if res == "fail" else "", BODY_PT,
               align=WD_ALIGN_PARAGRAPH.CENTER, color=RED if res == "fail" else None)
        _write_html(rr.cells[4], ln.get("method") or "", BODY_PT)

    # 本次稽核說明（缺失）＋ 本次稽核結果（統計）——版面同 sheets
    rr = add_row()
    box = _merge(rr, 0, 4)
    _write(box, "")
    _font(box.paragraphs[0].add_run("本次稽核說明:  "), BODY_PT)

    fails = [ln for ln in lines if ln.get("result") == "fail"]
    deficiency_text = "\n".join(f"{_item_no(ln)} {ln.get('name') or ''}".strip() for ln in fails)
    deficiency = box.add_table(rows=1, cols=2)
    deficiency.style = "Table Grid"
    deficiency.alignment = WD_TABLE_ALIGNMENT.CENTER
    _write(deficiency.rows[0].cells[0], "缺失", BODY_PT, align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
    _write(deficiency.rows[0].cells[1], deficiency_text, BODY_PT, color=RED)
    _fix_widths(deficiency, [Cm(2.6), Cm(14.0)])
    _row_height(deficiency.rows[0], 0.6)

    box.add_paragraph()
    _font(box.add_paragraph().add_run("本次稽核結果:"), RESULT_PT)

    stats = box.add_table(rows=1, cols=6)
    stats.style = "Table Grid"
    stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    score100: Optional[int] = sub.get("score100")
    failed = int(sub.get("fail_count") or 0)
    score_txt = "本月無" if score100 is None else f"{score100}分"
    vals = [("選定項數", str(int(sub.get("item_count") or 0))), ("不符合項數", str(failed)), ("本期稽核分數", score_txt)]
    for idx, (label, val) in enumerate(vals):
        if idx == 2:
            val_color = _score_color(None if score100 is None else score100 / 100)
        elif idx == 1 and failed > 0:
            val_color = RED
        else:
            val_color = None
        _write(stats.rows[0].cells[idx * 2], label, RESULT_PT, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(stats.rows[0].cells[idx * 2 + 1], val, RESULT_PT, align=WD_ALIGN_PARAGRAPH.CENTER, color=val_color)
        if idx == 1:
            for c_ in (stats.rows[0].cells[2], stats.rows[0].cells[3]):
                for run in c_.paragraphs[0].runs:
                    _highlight(run)
    _fix_widths(stats, [Cm(3.3), Cm(1.25), Cm(3.0), Cm(1.75), Cm(4.3), Cm(3.0)])
    _row_height(stats.rows[0], 0.6)
    box.add_paragraph()

    _fix_widths(t, COLS)

    sig = doc.add_table(rows=1, cols=3)
    sig.style = "Table Grid"
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    _row_height(sig.rows[0], 2.5)
    for idx, label in enumerate(["人資主管：", "受稽部門主管：", "稽核人員："]):
        cell = sig.rows[0].cells[idx]
        _write(cell, label, BODY_PT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if idx == 2 and auditor_name:
            _font(cell.add_paragraph().add_run(auditor_name), BODY_PT, bold=False)
    _fix_widths(sig, [Cm(5.75), Cm(5.25), Cm(7.0)])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
