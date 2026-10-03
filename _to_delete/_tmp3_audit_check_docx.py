"""
稽核檢查 — 單一部門「內部稽核檢查表」Word 匯出（2026-10-01）

格式比照使用者提供的「2026-..._202609-1 內部稽核檢查表.doc」：
  logo（維春商業開發）／標題／填表日期
  受稽單位｜文件編號、受稽人員｜稽核人員、文件名稱
  項次｜內部稽核要項｜稽核紀錄｜稽核結果｜建議
    1.  大項
    1.1 子項 ｜ 評語（判定色） ｜ V / X ｜ 建議（藍字）
    覆核. 8月待補正項目
        待補正內容 ｜ 覆核結果 ｜ V / X
  本次稽核說明：xx%（稽核子項數／達標項數／本期稽核分數、缺失）
  簽核列：執董／營運最高主管／受稽部門主管／稽核人員

只列出「這個部門有填評語或建議」的子項（沒填 ＝ 本期不查此項）；大項只在底下有子項時出現。
數字一律取 build_sheet_detail() 的系統計算結果，不在這裡重算。
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LOGO_PATH = Path(__file__).parent / "audit_check_assets" / "esse_logo.png"
FONT = "微軟正黑體"
# 顏色一律比照網頁稽核單（SheetEditor.tsx）：
#   缺失 #cf1322、建議 SUGGESTION_COLOR #1677ff、分數 100% 綠 #52c41a／未滿紅 #cf1322；
#   評語與稽核結果 → 判定類型主檔的顏色（使用者在「判定類型設定」自訂）
RED = RGBColor(0xCF, 0x13, 0x22)
GREEN = RGBColor(0x52, 0xC4, 0x1A)
BLUE = RGBColor(0x16, 0x77, 0xFF)


def _score_color(score: Optional[float]) -> Optional[RGBColor]:
    if score is None:
        return None
    return GREEN if score >= 1 else RED
SHADE = "F2F2F2"

# 主表欄寬（A4 直式，左右邊界 1.5cm → 可用 18cm）
COLS = [Cm(1.5), Cm(4.3), Cm(6.2), Cm(2.2), Cm(3.8)]


# ── 低階小工具 ────────────────────────────────────────────────────────────
def _font(run, size: float = 10, bold: bool = True, color: Optional[RGBColor] = None) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if color is not None:
        run.font.color.rgb = color


def _write(cell, text: str, size: float = 10, bold: bool = True,
           align=WD_ALIGN_PARAGRAPH.LEFT, color: Optional[RGBColor] = None) -> None:
    """把文字寫進格子（多行用換行分段），覆蓋原本內容。"""
    lines = (text or "").replace("\r\n", "\n").split("\n")
    first = cell.paragraphs[0]
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    for idx, line in enumerate(lines):
        p = first if idx == 0 else cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        _font(p.add_run(line), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _shade(cell, fill: str = SHADE) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _row_height(row, cm: float) -> None:
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trpr.append(h)


def _fix_widths(table, widths) -> None:
    """固定欄寬：tblGrid（LibreOffice／WPS 看這個）＋每格 tcW（Word 看這個）＋ fixed layout。"""
    table.autofit = False
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    grid = table._tbl.tblGrid
    for idx, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if idx < len(widths):
            gc.set(qn("w:w"), str(int(widths[idx].twips)))
    for row in table.rows:
        for idx, w in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = w


def _merge(row, start: int, end: int):
    return row.cells[start].merge(row.cells[end]) if end > start else row.cells[start]


def _symbol(code: Optional[str], types: Dict[str, dict]) -> str:
    """稽核結果欄：算達標 → V；不算達標 → X（與畫面判定一致）。"""
    if not code:
        return ""
    t = types.get(code)
    if t is None or t.get("counts_as_pass", True):
        return "V"
    return "X"


def _color(code: Optional[str], types: Dict[str, dict]) -> Optional[RGBColor]:
    """判定類型主檔的顏色（使用者在「判定類型設定」自訂，與網頁畫面同一來源）。"""
    if not code:
        return None
    t = types.get(code)
    hexv = (t or {}).get("color") or ""
    hexv = hexv.lstrip("#")
    if len(hexv) != 6:
        return None
    try:
        return RGBColor.from_string(hexv.upper())
    except ValueError:
        return None


def _month_label(period: Optional[str]) -> str:
    """'2026-08' → '8月'"""
    if period and len(period) >= 7 and period[4] == "-":
        try:
            return f"{int(period[5:7])}月"
        except ValueError:
            pass
    return "上期"


def _pct(score: Optional[float]) -> str:
    if score is None:
        return "本月無"
    return f"{round(score * 100)}%"


# ── 主函式 ────────────────────────────────────────────────────────────────
def build_inspection_docx(
    detail: dict,
    sheet_department_id: int,
    auditor_name: str,
    fill_date: str,
) -> bytes:
    """
    detail：audit_check_service.build_sheet_detail() 的結果
    fill_date：'YYYY.MM.DD'
    """
    dept = next(d for d in detail["departments"] if d["id"] == sheet_department_id)
    dept_index = [d["id"] for d in detail["departments"]].index(sheet_department_id) + 1
    company = detail.get("company_name") or ""
    period = detail.get("period") or ""
    types = {
        (t["code"] if isinstance(t, dict) else t.code): (
            t if isinstance(t, dict) else {"counts_as_pass": t.counts_as_pass, "color": t.color}
        )
        for t in detail.get("result_types", [])
    }

    cells = {
        c["sheet_item_id"]: c
        for c in detail["cells"]
        if c["sheet_department_id"] == sheet_department_id
        and ((c.get("comment") or "").strip() or (c.get("suggestion") or "").strip())
    }
    items: List[dict] = detail["items"]
    minors_by_major: Dict[int, List[dict]] = {}
    for it in items:
        if it["parent_sheet_item_id"] is not None and it["id"] in cells:
            minors_by_major.setdefault(it["parent_sheet_item_id"], []).append(it)
    majors = [it for it in items if it["parent_sheet_item_id"] is None and it["id"] in minors_by_major]

    review = next((r for r in detail.get("reviews", []) if r["sheet_department_id"] == sheet_department_id), None)
    score = next((s for s in detail.get("scores", []) if s["sheet_department_id"] == sheet_department_id), None)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.0)
    sec.bottom_margin = Cm(1.5)
    sec.header_distance = Cm(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)

    # ── 頁首：logo ──
    hp = sec.header.paragraphs[0]
    if LOGO_PATH.exists():
        hp.add_run().add_picture(str(LOGO_PATH), width=Cm(4.2))

    # ── 標題 / 填表日期 ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(p.add_run("內部稽核檢查表"), 16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(p.add_run(f"填表日期：{fill_date}"), 14)

    # ── 主表 ──
    t = doc.add_table(rows=0, cols=5)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    def add_row(height: float = 0.75):
        r = t.add_row()
        _row_height(r, height)
        return r

    r = add_row(0.9)
    _write(_merge(r, 0, 2), f"受稽單位：{company}", 14)
    _write(_merge(r, 3, 4), f"文件編號：{period.replace('-', '')}-{dept_index}", 14)

    r = add_row()
    _write(_merge(r, 0, 2), f"受稽人員：{dept['name']}", 12)
    _write(_merge(r, 3, 4), f"稽核人員：{auditor_name}", 12)

    r = add_row()
    _write(_merge(r, 0, 4), f"文件名稱：{company}系統建置作業-{dept['name']}", 12)

    r = add_row(1.0)
    for idx, label in enumerate(["項次", "內部稽核要項", "稽核紀錄", "稽核結果", "建議"]):
        _write(r.cells[idx], label, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
        _shade(r.cells[idx])

    def section_row(no: str, title: str, no_size: float = 14):
        rr = add_row(0.9)
        _write(rr.cells[0], no, no_size)
        _write(_merge(rr, 1, 4), title, 14)
        for c in rr.cells:
            _shade(c)

    for major in majors:
        section_row(f"{major['display_no']}." if major.get("display_no") else "", major["name"])
        for minor in minors_by_major[major["id"]]:
            c = cells[minor["id"]]
            rr = add_row(0.9)
            _write(rr.cells[0], minor.get("display_no") or "", 10, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _write(rr.cells[1], f"{minor['name']}{minor.get('scope_note') or ''}", 10)
            code = c.get("result_code") if (c.get("comment") or "").strip() else None
            color = _color(code, types)   # 評語與稽核結果依判定顏色（如扣分紅、建議藍）
            _write(rr.cells[2], c.get("comment") or "", 10, color=color)
            _write(rr.cells[3], _symbol(code, types), 11, align=WD_ALIGN_PARAGRAPH.CENTER, color=color)
            # 最右欄「建議」（原範本的「備註」欄，2026-10-03 使用者改名）：
            # 該格的建議，固定藍字 SUGGESTION_COLOR、與判定無關
            _write(rr.cells[4], (c.get("suggestion") or "").strip(), 10, color=BLUE)

    has_review = review and ((review.get("pending_text") or "").strip() or (review.get("result_text") or "").strip())
    if has_review:
        section_row("覆核.", f"{_month_label(review.get('source_period'))}待補正項目", 9)
        rr = add_row(1.2)
        _write(rr.cells[0], "")
        _write(rr.cells[1], review.get("pending_text") or "", 10,
               color=_color(review.get("pending_result"), types))
        _write(rr.cells[2], review.get("result_text") or "", 10,
               color=_color(review.get("result_status"), types))
        code = review.get("result_status") if (review.get("result_text") or "").strip() else review.get("pending_result")
        _write(rr.cells[3], _symbol(code, types), 11, align=WD_ALIGN_PARAGRAPH.CENTER, color=_color(code, types))
        _write(rr.cells[4], "")

    # ── 本次稽核說明 ──
    rr = add_row()
    box = _merge(rr, 0, 4)
    _write(box, "")
    p = box.paragraphs[0]
    _font(p.add_run("本次稽核說明:  "), 14)
    score_val = score["score"] if score else None
    _font(p.add_run(_pct(score_val)), 14, color=_score_color(score_val))

    inner = box.add_table(rows=2, cols=6)
    inner.style = "Table Grid"
    inner.alignment = WD_TABLE_ALIGNMENT.CENTER
    iw = [Cm(2.6), Cm(2.8), Cm(2.6), Cm(2.8), Cm(2.8), Cm(3.0)]
    sub = score["sub_count"] if score else 0
    passed = score["pass_count"] if score else 0
    vals = [("稽核子項數", str(sub)), ("達標項數", str(passed)), ("本期稽核分數", _pct(score["score"] if score else None))]
    for idx, (label, val) in enumerate(vals):
        _write(inner.rows[0].cells[idx * 2], label, 10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _write(inner.rows[0].cells[idx * 2 + 1], val, 10, align=WD_ALIGN_PARAGRAPH.CENTER,
               color=_score_color(score_val) if idx == 2 else None)
    _write(inner.rows[1].cells[0], "缺失", 10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RED)
    _write(inner.rows[1].cells[1].merge(inner.rows[1].cells[5]), dept.get("deficiency") or "", 10, color=RED)
    _fix_widths(inner, iw)
    for row in inner.rows:
        _row_height(row, 0.6)
    box.add_paragraph()   # 下方留白，比照原表

    _fix_widths(t, COLS)

    # ── 簽核列 ──
    sig = doc.add_table(rows=1, cols=4)
    sig.style = "Table Grid"
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    _row_height(sig.rows[0], 1.8)
    for idx, label in enumerate(["執董：", "營運最高主管：", "受稽部門主管：", "稽核人員："]):
        cell = sig.rows[0].cells[idx]
        _write(cell, label, 12)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        if idx == 3 and auditor_name:
            p = cell.add_paragraph()
            _font(p.add_run(auditor_name), 12, bold=False)
    _fix_widths(sig, [Cm(4.5)] * 4)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
